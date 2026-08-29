"""Render document outputs (Executive Summary, Advisory) into real .docx files.

Uses python-docx to produce a title plus labelled sections with bulleted
content (Key Findings, Key Facts, Recommendations, Action Items, etc.).
Returns the file bytes so the caller can persist them through the storage
abstraction.

Phase 8B hardening:
  * Rendering remains deterministic and LLM-free — it consumes the validated
    Pydantic models directly via the shared section normalizer and never
    invokes a language model.
  * Long titles / section lines are truncated defensively so a single
    pathological string cannot produce exploded run-on output.
  * Multi-line list items are normalized to a single line per paragraph so
    layout stays deterministic.
  * Empty optional collections are skipped safely.
  * ``parse_docx()`` round-trip verification is preserved.
  * Only the bundled ``python-docx`` library is used (no new dependencies).
"""

from __future__ import annotations

from io import BytesIO

from docx import Document

from app.transformation.output_schemas import Advisory, ExecutiveSummary
from app.transformation.render._documents import (
    advisory_sections,
    executive_summary_sections,
)

DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def render_executive_summary_docx(summary: ExecutiveSummary) -> bytes:
    """Render an ExecutiveSummary into DOCX bytes."""
    return _render_docx(executive_summary_sections(summary))


def render_advisory_docx(advisory: Advisory) -> bytes:
    """Render an Advisory into DOCX bytes."""
    return _render_docx(advisory_sections(advisory))


def _render_docx(sections: list[tuple[str | None, list[str]]]) -> bytes:
    """Render normalized sections into a deterministic DOCX document."""
    document = Document()

    for index, (heading, lines) in enumerate(sections):
        if index == 0:
            # Leading section carries the title with a None heading.
            for title_line in lines:
                document.add_heading(title_line, level=0)
            continue

        if heading is not None:
            document.add_heading(heading, level=1)
        for line in lines:
            document.add_paragraph(line, style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def parse_docx(byte_data: bytes) -> Document:
    """Parse DOCX bytes back into a python-docx Document (verification)."""
    return Document(BytesIO(byte_data))


__all__ = [
    "DOCX_MIME_TYPE",
    "parse_docx",
    "render_advisory_docx",
    "render_executive_summary_docx",
]
