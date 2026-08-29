"""Render a validated PresentationStructure into a real .pptx file.

Uses python-pptx to produce slides with titles, key-message content,
supporting points, speaker notes and (on the notes slide) a visual
recommendation line.  Returns the file bytes so the caller can persist them
through the storage abstraction.
"""

from __future__ import annotations

from io import BytesIO

from pptx import Presentation

from app.transformation.output_schemas import PresentationStructure


PPTX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


def render_presentation(structure: PresentationStructure) -> bytes:
    """Render a presentation structure into PPTX bytes."""
    if not structure.slides:
        raise ValueError("Presentation requires at least one slide.")

    prs = Presentation()
    # Use the title+content layout (index 1) which exposes title + body.
    for slide in structure.slides:
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        s = prs.slides.add_slide(layout)
        if s.shapes.title is not None:
            s.shapes.title.text = slide.title

        body_lines: list[str] = []
        if slide.key_message:
            body_lines.append(slide.key_message)
        for point in slide.supporting_points:
            body_lines.append(f"• {point}")
        if slide.visual_recommendation and len(prs.slide_layouts) > 1:
            # Only append visual rec to the body if we have a real content layout.
            body_lines.append(f"[Visual: {slide.visual_recommendation}]")

        body_placeholder = None
        for placeholder in s.placeholders:
            if placeholder.placeholder_format.idx == 1:
                body_placeholder = placeholder
                break
        if body_placeholder is not None and body_lines:
            tf = body_placeholder.text_frame
            tf.clear()
            for index, line in enumerate(body_lines):
                paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
                paragraph.text = line

        if slide.speaker_notes or slide.visual_recommendation:
            notes_frame = s.notes_slide.notes_text_frame
            note_parts = list(slide.speaker_notes)
            if slide.visual_recommendation:
                note_parts.append(f"Visual recommendation: {slide.visual_recommendation}")
            notes_frame.text = "\n".join(note_parts)

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def parse_pptx(byte_data: bytes) -> Presentation:
    """Parse PPTX bytes back into a python-pptx Presentation (verification)."""
    return Presentation(BytesIO(byte_data))
